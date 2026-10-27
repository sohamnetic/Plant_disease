"""Every image used to build the app's two models, as one Excel workbook.

  Classifier  experiments/experiment_006_v3_measured_severity (EfficientNet-B0 x5)
              trained on splits_v3/train_pool.csv (= its folds.csv, 5,764 images).
              Its calibration / 'not sure' layer was fitted on the same images'
              out-of-fold predictions, so it adds no images.
  YOLO        experiments/seg_001_yolo11s trained on the 299 outlined photos with
              seg_split == "train" in outputs/seg/labelled_images.csv (copied to
              data_external/seg_yolo/images/train); all 299 are also in the classifier pool.
  Test only   splits_v3/external_test.csv, splits/test.csv, and the 73 YOLO held-out
              photos (taken from those two tests): never trained on by either model.

dataset_balanced/ is not listed: no model was trained on it.

  python src/list_training_images.py
Output: outputs/final/TRAINING_IMAGES_LIST.xlsx
"""

from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import config as cfg

ROOT = cfg.PROJECT_ROOT
OUT = cfg.OUTPUTS_DIR / "final" / "TRAINING_IMAGES_LIST.xlsx"
YOLO_TRAIN_DIR = ROOT / "data_external" / "seg_yolo" / "images" / "train"

CLASS = {"bacterial_spot_rot": "Bacterial spot rot", "black_rot": "Black rot",
         "downy_mildew": "Downy mildew", "no_disease": "Healthy"}
SOURCE = {"vegnet": "VegNet (our original dataset)",
          "c2h": "Cauliflower Diseases Identification 2025 (c2h)",
          "sppu": "SPPU Cauliflower",
          "cleaf": "Cauliflower Leaf Diseases 2025 (cleaf)",
          "vegleaf": "Vegetable Leaf Collection (cauliflower leaves)",
          "healthy_leaf_crop": "Healthy-leaf crops (made by us from VegNet photos)"}
KIND = {"original": "Original photo", "augmented": "Augmented copy (from the VegNet publisher)",
        "external": "Outside photo", "bg_removed_copy": "VegNet photo with background removed (published in c2h)",
        "derived": "Leaf crop cut from a VegNet healthy photo"}
CLASSES = list(CLASS.values())

FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=14, bold=True, color="2F6B45")
HEAD_FILL = PatternFill("solid", fgColor="E3EDE6")
TOTAL_FILL = PatternFill("solid", fgColor="F5F8F6")
THIN = Border(bottom=Side(style="thin", color="C9D3CD"))


def nice(df: pd.DataFrame) -> pd.DataFrame:
    return df.assign(cls=df["class"].map(CLASS), src=df["source"].map(SOURCE),
                     kind=df["source_type"].map(KIND) if "source_type" in df else "")


def list_sheet(wb: Workbook, name: str, headers: list[str], rows: list[list], widths: list[int]) -> None:
    ws = wb.create_sheet(name)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill, c.border = BOLD, HEAD_FILL, THIN
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def count_table(ws, top: int, title: str, sheet: str, row_col: str, row_labels: list[str],
                extra: tuple[str, str] | None = None) -> int:
    """Rows = row_labels matched in column row_col of `sheet`, columns = classes (column D there).
    extra = (column, value) adds one more COUNTIFS condition. Returns the next free row."""
    ws.cell(top, 1, title).font = BOLD
    hdr = top + 1
    for j, h in enumerate(["", *CLASSES, "Total"], 1):
        c = ws.cell(hdr, j, h)
        c.font, c.fill, c.border = BOLD, HEAD_FILL, THIN
    q = f"'{sheet}'!"
    for i, lab in enumerate(row_labels):
        r = hdr + 1 + i
        ws.cell(r, 1, lab).font = FONT
        for j in range(len(CLASSES)):
            col = get_column_letter(2 + j)
            cond = f"{q}${row_col}:${row_col},$A{r},{q}$D:$D,{col}${hdr}"
            if extra:
                cond += f",{q}${extra[0]}:${extra[0]},\"{extra[1]}\""
            ws.cell(r, 2 + j, f"=COUNTIFS({cond})").font = FONT
        ws.cell(r, 2 + len(CLASSES), f"=SUM(B{r}:{get_column_letter(1 + len(CLASSES))}{r})").font = BOLD
    tot = hdr + 1 + len(row_labels)
    ws.cell(tot, 1, "Total").font = BOLD
    for j in range(len(CLASSES) + 1):
        col = get_column_letter(2 + j)
        c = ws.cell(tot, 2 + j, f"=SUM({col}{hdr + 1}:{col}{tot - 1})")
        c.font = BOLD
    for j in range(1, len(CLASSES) + 3):
        ws.cell(tot, j).fill = TOTAL_FILL
    return tot + 2


def main() -> None:
    pool = nice(pd.read_csv(ROOT / "experiments/experiment_006_v3_measured_severity/folds.csv"))
    li = pd.read_csv(cfg.OUTPUTS_DIR / "seg" / "labelled_images.csv")
    yolo = nice(li[li["seg_split"] == "train"].merge(pool[["filepath", "source_type"]], on="filepath"))
    yolo["copy"] = [f"{i:04d}_{s}_{c}.jpg" for i, s, c in zip(yolo["idx"], yolo["source"], yolo["class"])]
    missing = [f for f in yolo["copy"] if not (YOLO_TRAIN_DIR / f).exists()]
    assert not missing and len(yolo) == len(list(YOLO_TRAIN_DIR.glob("*.jpg"))), missing[:3]
    ext = nice(pd.read_csv(ROOT / "splits_v3" / "external_test.csv")).assign(test="Outside photos test (external)")
    vt = nice(pd.read_csv(cfg.SPLITS_DIR / "test.csv").assign(source="vegnet")).assign(test="VegNet test split")
    seg_eval = set(li[li["seg_split"] == "eval"]["filepath"])
    tests = pd.concat([ext, vt], ignore_index=True)
    assert not tests["filepath"].isin(pool["filepath"]).any()

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Images used to build the cauliflower disease models"
    ws["A1"].font = TITLE
    notes = [
        "Disease classifier (EfficientNet-B0, 5 models): trained on every image in sheet 'Classifier training'. "
        "Each image was used by 4 of the 5 models; the 'Validation fold' column says which model checked itself on it instead.",
        "Its confidence / 'not sure' layer was tuned on those same images (out-of-fold predictions), so it uses no extra images.",
        "Lesion segmentation (YOLO11s-seg): trained on the photos in sheet 'YOLO training', each outlined by hand "
        "(leaf, head, lesion). All of them are also in the classifier list.",
        "Sheet 'Test only' lists the photos used only to measure the models. Neither model was ever trained on them.",
        "Not listed: dataset_balanced/ (originals + our own augmented copies, built 12 Oct 2026). No model was trained on it.",
    ]
    for i, t in enumerate(notes):
        c = ws.cell(3 + i, 1, t)
        c.font = FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=3 + i, start_column=1, end_row=3 + i, end_column=6)
        ws.row_dimensions[3 + i].height = 28
    r = 3 + len(notes) + 1
    r = count_table(ws, r, "Classifier training images, by source", "Classifier training", "E", list(SOURCE.values()))
    r = count_table(ws, r, "Classifier training images, by type", "Classifier training", "F", list(KIND.values()))
    r = count_table(ws, r, "YOLO training images, by source", "YOLO training", "E",
                    [s for s in SOURCE.values() if (yolo["src"] == s).any()])
    r = count_table(ws, r, "Test-only images (never trained on)", "Test only", "E", ["Outside photos test (external)",
                    "VegNet test split"])
    count_table(ws, r, "  of which used to test YOLO (marked 'Yes' in sheet 'Test only')", "Test only", "E",
                ["Outside photos test (external)", "VegNet test split"], extra=("G", "Yes"))
    ws.column_dimensions["A"].width = 58
    for col in "BCDEF":
        ws.column_dimensions[col].width = 17

    def name(p: str) -> str:
        return Path(p).name

    list_sheet(wb, "Classifier training",
               ["No.", "File name", "Full path (inside D:\\project-1)", "Disease", "Source dataset", "Image type",
                "Validation fold", "Severity grade (0-4)", "Also used for YOLO"],
               [[i + 1, name(r.filepath), r.filepath, r.cls, r.src, r.kind, int(r.fold),
                 int(r.grade) if r.sev_kind == 0 else "not graded (diseased)",
                 "Yes" if r.filepath in set(yolo["filepath"]) else ""]
                for i, r in enumerate(pool.sort_values(["class", "source", "source_type", "filepath"]).itertuples())],
               [7, 40, 95, 18, 46, 52, 15, 19, 18])
    list_sheet(wb, "YOLO training",
               ["No.", "File name", "Full path (inside D:\\project-1)", "Disease", "Source dataset", "Image type",
                "Severity grade (0-4)", "Copy used by YOLO (data_external/seg_yolo/images/train)"],
               [[i + 1, name(r.filepath), r.filepath, r.cls, r.src, r.kind,
                 "" if pd.isna(r.grade) else int(r.grade), r.copy]
                for i, r in enumerate(yolo.sort_values(["class", "source", "filepath"]).itertuples())],
               [7, 40, 95, 18, 46, 52, 19, 50])
    list_sheet(wb, "Test only",
               ["No.", "File name", "Full path (inside D:\\project-1)", "Disease", "Test set", "Source dataset",
                "Also used to test YOLO", "Image type"],
               [[i + 1, name(r.filepath), r.filepath, r.cls, r.test, r.src, "Yes" if r.filepath in seg_eval else "",
                 r.kind] for i, r in enumerate(tests.sort_values(["test", "class", "source", "filepath"]).itertuples())],
               [7, 40, 95, 18, 30, 46, 22, 52])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("written", OUT, "| classifier", len(pool), "| yolo", len(yolo), "| test only", len(tests),
          "| yolo test", int(tests["filepath"].isin(seg_eval).sum()))


if __name__ == "__main__":
    main()
